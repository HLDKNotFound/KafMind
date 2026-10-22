import asyncio
import json
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
import uuid

from aiokafka import AIOKafkaConsumer
from langchain_text_splitters import RecursiveCharacterTextSplitter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.database import async_session_maker
from app.core.qdrant import get_qdrant_client
from app.core.redis import get_redis_client
from app.models.file import DocumentFile, FileStatus
from app.services.qdrant_service import upsert_document_chunks

logger = logging.getLogger(__name__)
settings = get_settings()


def extract_text_from_pdf(file_path: str) -> str:
    """Extract plain text from PDF using pypdf."""
    from pypdf import PdfReader
    reader = PdfReader(file_path)
    pages_text = []
    for page in reader.pages:
        extracted = page.extract_text()
        if extracted:
            pages_text.append(extracted)
    return "\n\n".join(pages_text).strip()


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from DOCX paragraphs and tables using python-docx."""
    import docx
    doc = docx.Document(file_path)
    lines: List[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_str:
                lines.append(row_str)
    return "\n\n".join(lines).strip()


def extract_text_from_txt(file_path: str) -> str:
    """Extract text from plain text or markdown file with UTF-8 encoding."""
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        return f.read().strip()


def extract_text(file_path: str) -> str:
    """Detect file type and extract full text content."""
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"File not found on disk: {file_path}")

    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext in {".txt", ".md"}:
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file format for ingestion: '{ext}'")


def chunk_document_text(
    text: str,
    chunk_size: int = 800,
    chunk_overlap: int = 150,
) -> List[str]:
    """Split text into overlapping chunks using RecursiveCharacterTextSplitter."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", " ", ""],
    )
    return splitter.split_text(text)


async def publish_file_status(
    redis_client: Any,
    user_id: Union[str, uuid.UUID],
    file_id: Union[str, uuid.UUID],
    status: str,
    chunk_count: int = 0,
    error: Optional[str] = None,
) -> None:
    """Broadcast real-time file processing status to Redis Pub/Sub."""
    channel = f"file_status:{user_id}"
    payload = {
        "file_id": str(file_id),
        "status": status,
        "chunk_count": chunk_count,
        "error": error,
    }
    try:
        await redis_client.publish(channel, json.dumps(payload))
        logger.info(f"Published status to {channel}: {payload}")
    except Exception as e:
        logger.warning(f"Failed to publish status to {channel}: {e}")


async def process_file_event(
    event: Dict[str, Any],
    db_session: Optional[AsyncSession] = None,
    redis_client: Optional[Any] = None,
    qdrant_client: Optional[Any] = None,
) -> None:
    """Process a single file.uploaded event end-to-end: parse, chunk, embed, index, and notify."""
    file_id_str = event.get("file_id")
    user_id_str = event.get("user_id")
    file_path = event.get("file_path")
    filename = event.get("filename")

    if not file_id_str or not user_id_str or not file_path:
        logger.error(f"Invalid file.uploaded event received: {event}")
        return

    file_uuid = uuid.UUID(file_id_str)
    user_uuid = uuid.UUID(user_id_str)

    if redis_client is None:
        redis_client = get_redis_client()
    if qdrant_client is None:
        qdrant_client = get_qdrant_client()

    # Determine if we manage the DB session lifecycle or use an injected one
    managed_session = db_session is None
    session = db_session if not managed_session else async_session_maker()

    try:
        # 1. Fetch file record from DB
        stmt = select(DocumentFile).where(DocumentFile.id == file_uuid)
        res = await session.execute(stmt)
        doc_file = res.scalar_one_or_none()
        if doc_file is None:
            logger.error(f"DocumentFile record {file_uuid} not found in database.")
            return

        # 2. Update status to PROCESSING and broadcast
        doc_file.status = FileStatus.PROCESSING.value
        await session.commit()
        await publish_file_status(
            redis_client=redis_client,
            user_id=user_uuid,
            file_id=file_uuid,
            status=FileStatus.PROCESSING.value,
            chunk_count=0,
        )

        # 3. Extract text off-thread
        text = await asyncio.to_thread(extract_text, file_path)
        if not text or not text.strip():
            raise ValueError(f"Extracted content is empty for file: '{filename}'")

        # 4. Chunk text
        chunks = chunk_document_text(text)
        if not chunks:
            raise ValueError(f"No chunks generated for file: '{filename}'")

        # 5. Embed and upsert to Qdrant
        await upsert_document_chunks(
            user_id=user_uuid,
            file_id=file_uuid,
            filename=filename or doc_file.filename,
            chunks=chunks,
            client=qdrant_client,
        )

        # 6. Update DB record to READY
        doc_file.status = FileStatus.READY.value
        doc_file.chunk_count = len(chunks)
        doc_file.error_message = None
        await session.commit()

        # 7. Broadcast READY status via Redis Pub/Sub
        await publish_file_status(
            redis_client=redis_client,
            user_id=user_uuid,
            file_id=file_uuid,
            status=FileStatus.READY.value,
            chunk_count=len(chunks),
        )
        logger.info(f"Successfully indexed file_id={file_uuid} with {len(chunks)} chunks.")

    except Exception as e:
        logger.error(f"Error ingesting file_id={file_uuid}: {e}", exc_info=True)
        try:
            # Mark file as FAILED in DB
            stmt = select(DocumentFile).where(DocumentFile.id == file_uuid)
            res = await session.execute(stmt)
            doc_file = res.scalar_one_or_none()
            if doc_file:
                doc_file.status = FileStatus.FAILED.value
                doc_file.error_message = str(e)
                await session.commit()

            # Broadcast FAILED status
            await publish_file_status(
                redis_client=redis_client,
                user_id=user_uuid,
                file_id=file_uuid,
                status=FileStatus.FAILED.value,
                chunk_count=0,
                error=str(e),
            )
        except Exception as db_err:
            logger.error(f"Failed to record failure state for file_id={file_uuid}: {db_err}")
    finally:
        if managed_session:
            await session.close()


async def run_ingestion_worker() -> None:
    """Kafka consumer loop consuming file.uploaded topic."""
    consumer = AIOKafkaConsumer(
        settings.KAFKA_TOPIC_FILE_UPLOADED,
        bootstrap_servers=settings.KAFKA_BOOTSTRAP_SERVERS,
        group_id=settings.KAFKA_CONSUMER_GROUP,
        value_deserializer=lambda v: json.loads(v.decode("utf-8")),
        auto_offset_reset="earliest",
        enable_auto_commit=True,
    )
    await consumer.start()
    logger.info(
        f"Kafka Ingestion Worker started, listening to '{settings.KAFKA_TOPIC_FILE_UPLOADED}' "
        f"on group '{settings.KAFKA_CONSUMER_GROUP}'."
    )

    try:
        async for msg in consumer:
            logger.info(f"Received message: offset={msg.offset}, key={msg.key}")
            try:
                event_data = msg.value
                await process_file_event(event_data)
            except Exception as e:
                logger.error(f"Failed to process message {msg.offset}: {e}", exc_info=True)
    finally:
        await consumer.stop()
        logger.info("Kafka Ingestion Worker stopped.")


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    asyncio.run(run_ingestion_worker())
