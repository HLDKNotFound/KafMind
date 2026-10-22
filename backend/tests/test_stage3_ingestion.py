import asyncio
import io
import json
from pathlib import Path
from typing import AsyncGenerator, Dict, List
from unittest.mock import AsyncMock, patch
import uuid

import docx
import httpx
import pytest
import pytest_asyncio
from qdrant_client import AsyncQdrantClient, models
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from app.core.config import get_settings
from app.core.database import Base, get_db
from app.main import app
from app.models.file import DocumentFile, FileStatus
from app.models.user import User
from app.services.qdrant_service import build_user_filter, init_qdrant_collection
from app.workers.ingestion_worker import process_file_event

settings = get_settings()


class InMemoryRedisMock:
    """Async in-memory key-value store and Pub/Sub simulating Redis."""

    def __init__(self):
        self.store: Dict[str, tuple] = {}
        self.channels: Dict[str, List[asyncio.Queue]] = {}

    async def setex(self, key: str, time_secs: int, value: str) -> bool:
        self.store[key] = (value, asyncio.get_event_loop().time() + time_secs)
        return True

    async def exists(self, key: str) -> int:
        if key in self.store:
            _, exp = self.store[key]
            if asyncio.get_event_loop().time() < exp:
                return 1
            del self.store[key]
        return 0

    async def ping(self) -> bool:
        return True

    async def publish(self, channel: str, message: str) -> int:
        subs = self.channels.get(channel, [])
        for q in subs:
            await q.put({"type": "message", "channel": channel, "data": message})
        return len(subs)

    def pubsub(self):
        return MockPubSub(self)

    async def aclose(self) -> None:
        self.store.clear()
        self.channels.clear()


class MockPubSub:
    """Mock Redis PubSub subscriber object with async queues."""

    def __init__(self, redis_mock: InMemoryRedisMock):
        self.redis = redis_mock
        self.subscribed_channels = set()
        self.queue: asyncio.Queue = asyncio.Queue()

    async def subscribe(self, *channels: str):
        for ch in channels:
            self.subscribed_channels.add(ch)
            if ch not in self.redis.channels:
                self.redis.channels[ch] = []
            self.redis.channels[ch].append(self.queue)

    async def unsubscribe(self, *channels: str):
        for ch in channels:
            self.subscribed_channels.discard(ch)
            if ch in self.redis.channels and self.queue in self.redis.channels[ch]:
                self.redis.channels[ch].remove(self.queue)

    async def get_message(self, ignore_subscribe_messages: bool = True, timeout: float = 0.5):
        try:
            return await asyncio.wait_for(self.queue.get(), timeout=timeout)
        except asyncio.TimeoutError:
            return None

    async def aclose(self):
        for ch in list(self.subscribed_channels):
            await self.unsubscribe(ch)


@pytest_asyncio.fixture(scope="function")
async def test_env(tmp_path: Path):
    """Provide isolated SQLite database, in-memory Qdrant client, and Redis mock."""
    # 1. SQLite in-memory engine
    engine = create_async_engine(
        "sqlite+aiosqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    session_maker = async_sessionmaker(
        bind=engine,
        class_=AsyncSession,
        expire_on_commit=False,
        autocommit=False,
        autoflush=False,
    )

    async def override_get_db() -> AsyncGenerator[AsyncSession, None]:
        async with session_maker() as session:
            try:
                yield session
            finally:
                await session.close()

    app.dependency_overrides[get_db] = override_get_db

    # 2. Redis Mock
    mock_redis = InMemoryRedisMock()

    # 3. Real in-memory Qdrant client
    qdrant_mem_client = AsyncQdrantClient(":memory:")
    await init_qdrant_collection(qdrant_mem_client)

    # 4. Mock Kafka Producer
    published_events = []

    async def mock_publish(file_id, user_id, file_path, filename, producer=None):
        evt = {
            "file_id": str(file_id),
            "user_id": str(user_id),
            "file_path": str(file_path),
            "filename": filename,
        }
        published_events.append(evt)
        return evt

    upload_folder = tmp_path / "uploads"
    upload_folder.mkdir()

    with (
        patch("app.core.security.get_redis_client", return_value=mock_redis),
        patch("app.api.v1.files.get_redis_client", return_value=mock_redis),
        patch("app.core.qdrant.get_qdrant_client", return_value=qdrant_mem_client),
        patch("app.services.qdrant_service.get_qdrant_client", return_value=qdrant_mem_client),
        patch("app.api.v1.files.publish_file_uploaded", side_effect=mock_publish),
        patch.object(settings, "UPLOAD_DIR", str(upload_folder)),
    ):
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app),
            base_url="http://test",
        ) as client:
            yield {
                "client": client,
                "session_maker": session_maker,
                "redis": mock_redis,
                "qdrant": qdrant_mem_client,
                "published_events": published_events,
                "upload_dir": upload_folder,
            }

    app.dependency_overrides.clear()
    await qdrant_mem_client.close()
    await engine.dispose()


async def register_and_login(client: httpx.AsyncClient, email: str = "author@example.com") -> tuple[str, str]:
    """Helper to register and login a user, returning access_token and user_id."""
    res = await client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "password123", "full_name": "Author User"},
    )
    user_id = res.json()["id"]

    login_res = await client.post(
        "/api/v1/auth/login",
        json={"email": email, "password": "password123"},
    )
    token = login_res.json()["access_token"]
    return token, user_id


@pytest.mark.asyncio
async def test_file_upload_202_accepted(test_env: dict):
    """Test POST /api/v1/files/upload stores file, creates PENDING record, publishes Kafka event, and returns 202."""
    client: httpx.AsyncClient = test_env["client"]
    token, user_id = await register_and_login(client, "uploader@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    file_content = b"KafMind is an Agentic RAG system built with FastAPI, Kafka, and Qdrant."
    files = {"file": ("architecture.txt", file_content, "text/plain")}

    response = await client.post("/api/v1/files/upload", files=files, headers=headers)
    assert response.status_code == 202

    data = response.json()
    assert data["message"] == "File upload accepted and queued for processing"
    assert "file_id" in data
    file_id = data["file_id"]

    file_meta = data["file"]
    assert file_meta["filename"] == "architecture.txt"
    assert file_meta["status"] == "PENDING"
    assert file_meta["chunk_count"] == 0
    assert file_meta["file_size"] == len(file_content)

    # Verify physical file exists on disk
    saved_path = Path(file_meta["file_path"])
    assert saved_path.exists()
    assert saved_path.read_bytes() == file_content

    # Verify Kafka event published
    assert len(test_env["published_events"]) == 1
    event = test_env["published_events"][0]
    assert event["file_id"] == file_id
    assert event["user_id"] == user_id
    assert event["filename"] == "architecture.txt"


@pytest.mark.asyncio
async def test_file_upload_unsupported_format_rejected(test_env: dict):
    """Test POST /api/v1/files/upload rejects non-allowed file types with 400 Bad Request."""
    client: httpx.AsyncClient = test_env["client"]
    token, _ = await register_and_login(client, "badfile@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    files = {"file": ("malicious.exe", b"MZ...", "application/x-msdownload")}
    response = await client.post("/api/v1/files/upload", files=files, headers=headers)
    assert response.status_code == 400
    assert "Unsupported file format" in response.json()["detail"]


@pytest.mark.asyncio
async def test_worker_ingestion_and_qdrant_payload(test_env: dict):
    """Test Ingestion Worker parses text, chunks, embeds with FastEmbed, and upserts vectors with multi-tenant payload."""
    client: httpx.AsyncClient = test_env["client"]
    session_maker = test_env["session_maker"]
    qdrant: AsyncQdrantClient = test_env["qdrant"]
    redis: InMemoryRedisMock = test_env["redis"]

    token, user_id = await register_and_login(client, "ingest_user@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    # Subscribe to Redis Pub/Sub for status monitoring
    pubsub = redis.pubsub()
    await pubsub.subscribe(f"file_status:{user_id}")

    # Upload document
    doc_text = (
        "Antigravity is an advanced coding agent by Google DeepMind. "
        "It supports rich async pipelines, vector embeddings, and multi-tenant security.\n\n"
        "Apache Kafka enables resilient and horizontally scalable event streaming."
    )
    files = {"file": ("antigravity_guide.md", doc_text.encode("utf-8"), "text/markdown")}
    upload_res = await client.post("/api/v1/files/upload", files=files, headers=headers)
    assert upload_res.status_code == 202
    file_id = upload_res.json()["file_id"]
    event = test_env["published_events"][0]

    # Execute Worker on the event
    async with session_maker() as session:
        await process_file_event(
            event=event,
            db_session=session,
            redis_client=redis,
            qdrant_client=qdrant,
        )

    # Verify DB record updated to READY with chunk_count > 0
    async with session_maker() as session:
        res = await session.execute(select(DocumentFile).where(DocumentFile.id == uuid.UUID(file_id)))
        doc = res.scalar_one()
        assert doc.status == FileStatus.READY.value
        assert doc.chunk_count > 0
        assert doc.error_message is None

    # Verify Qdrant points indexed with multi-tenant payload
    filter_condition = build_user_filter(user_id=user_id)
    scroll_res, _ = await qdrant.scroll(
        collection_name=settings.QDRANT_COLLECTION_NAME,
        scroll_filter=filter_condition,
        with_payload=True,
        with_vectors=True,
    )
    assert len(scroll_res) == doc.chunk_count
    for point in scroll_res:
        payload = point.payload
        assert payload["user_id"] == str(user_id)
        assert payload["file_id"] == str(file_id)
        assert payload["filename"] == "antigravity_guide.md"
        assert len(payload["text"]) > 0
        assert "chunk_index" in payload
        # Verify 384-dimensional FastEmbed dense vector
        assert len(point.vector) == settings.EMBEDDING_DIMENSION

    # Verify Redis status messages received (PROCESSING then READY)
    msg1 = await pubsub.get_message(timeout=1.0)
    assert msg1 is not None
    data1 = json.loads(msg1["data"])
    assert data1["status"] == "PROCESSING"
    assert data1["file_id"] == file_id

    msg2 = await pubsub.get_message(timeout=1.0)
    assert msg2 is not None
    data2 = json.loads(msg2["data"])
    assert data2["status"] == "READY"
    assert data2["chunk_count"] == doc.chunk_count
    await pubsub.aclose()


@pytest.mark.asyncio
async def test_get_files_multi_tenant_isolation(test_env: dict):
    """Test GET /api/v1/files enforces strict multi-tenant isolation."""
    client: httpx.AsyncClient = test_env["client"]

    # User 1
    token1, _ = await register_and_login(client, "alice@example.com")
    await client.post(
        "/api/v1/files/upload",
        files={"file": ("alice_notes.txt", b"Alice confidential notes", "text/plain")},
        headers={"Authorization": f"Bearer {token1}"},
    )

    # User 2
    token2, _ = await register_and_login(client, "bob@example.com")
    await client.post(
        "/api/v1/files/upload",
        files={"file": ("bob_notes.txt", b"Bob private document", "text/plain")},
        headers={"Authorization": f"Bearer {token2}"},
    )

    # Alice lists files
    alice_res = await client.get("/api/v1/files", headers={"Authorization": f"Bearer {token1}"})
    assert alice_res.status_code == 200
    alice_files = alice_res.json()
    assert len(alice_files) == 1
    assert alice_files[0]["filename"] == "alice_notes.txt"

    # Bob lists files
    bob_res = await client.get("/api/v1/files", headers={"Authorization": f"Bearer {token2}"})
    assert bob_res.status_code == 200
    bob_files = bob_res.json()
    assert len(bob_files) == 1
    assert bob_files[0]["filename"] == "bob_notes.txt"


@pytest.mark.asyncio
async def test_delete_file_and_qdrant_cleanup(test_env: dict):
    """Test DELETE /api/v1/files/{file_id} removes DB record, Qdrant vectors, and local file."""
    client: httpx.AsyncClient = test_env["client"]
    session_maker = test_env["session_maker"]
    qdrant: AsyncQdrantClient = test_env["qdrant"]
    redis = test_env["redis"]

    token, user_id = await register_and_login(client, "deleter@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    upload_res = await client.post(
        "/api/v1/files/upload",
        files={"file": ("to_delete.txt", b"This file will be deleted.", "text/plain")},
        headers=headers,
    )
    file_id = upload_res.json()["file_id"]
    saved_path = Path(upload_res.json()["file"]["file_path"])
    event = test_env["published_events"][0]

    # Process through worker
    async with session_maker() as session:
        await process_file_event(event=event, db_session=session, redis_client=redis, qdrant_client=qdrant)

    # Verify vectors indexed
    res_before, _ = await qdrant.scroll(
        collection_name=settings.QDRANT_COLLECTION_NAME,
        scroll_filter=build_user_filter(user_id=user_id),
    )
    assert len(res_before) > 0
    assert saved_path.exists()

    # Call DELETE endpoint
    del_res = await client.delete(f"/api/v1/files/{file_id}", headers=headers)
    assert del_res.status_code == 200
    assert del_res.json()["message"] == "File deleted successfully"

    # Verify vectors removed from Qdrant
    res_after, _ = await qdrant.scroll(
        collection_name=settings.QDRANT_COLLECTION_NAME,
        scroll_filter=build_user_filter(user_id=user_id),
    )
    assert len(res_after) == 0

    # Verify file removed from disk
    assert not saved_path.exists()

    # Verify DB record gone
    async with session_maker() as session:
        db_res = await session.execute(select(DocumentFile).where(DocumentFile.id == uuid.UUID(file_id)))
        assert db_res.scalar_one_or_none() is None


@pytest.mark.asyncio
async def test_worker_ingestion_docx(test_env: dict, tmp_path: Path):
    """Test Ingestion Worker correctly parses and indexes DOCX documents."""
    client: httpx.AsyncClient = test_env["client"]
    session_maker = test_env["session_maker"]
    qdrant: AsyncQdrantClient = test_env["qdrant"]
    redis = test_env["redis"]

    token, user_id = await register_and_login(client, "docx_user@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    # Create real DOCX file
    docx_file = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_heading("KafMind Architecture Document", level=1)
    doc.add_paragraph("This document tests ingestion of Microsoft Word docx format.")
    doc.save(docx_file)

    with open(docx_file, "rb") as f:
        upload_res = await client.post(
            "/api/v1/files/upload",
            files={"file": ("sample.docx", f.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers=headers,
        )
    assert upload_res.status_code == 202
    file_id = upload_res.json()["file_id"]
    event = test_env["published_events"][0]

    # Run worker
    async with session_maker() as session:
        await process_file_event(event=event, db_session=session, redis_client=redis, qdrant_client=qdrant)

    # Verify READY
    async with session_maker() as session:
        res = await session.execute(select(DocumentFile).where(DocumentFile.id == uuid.UUID(file_id)))
        doc_record = res.scalar_one()
        assert doc_record.status == FileStatus.READY.value
        assert doc_record.chunk_count > 0


@pytest.mark.asyncio
async def test_worker_corrupt_file_handling(test_env: dict, tmp_path: Path):
    """Test Ingestion Worker gracefully marks corrupt or empty documents as FAILED."""
    client: httpx.AsyncClient = test_env["client"]
    session_maker = test_env["session_maker"]
    qdrant: AsyncQdrantClient = test_env["qdrant"]
    redis: InMemoryRedisMock = test_env["redis"]

    token, user_id = await register_and_login(client, "corrupt_user@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    # Upload empty file
    empty_file = tmp_path / "empty.txt"
    empty_file.write_text("   \n\t  ")  # Only whitespace

    # Directly create file record and event to bypass endpoint validation
    file_id = uuid.uuid4()
    async with session_maker() as session:
        record = DocumentFile(
            id=file_id,
            user_id=uuid.UUID(user_id),
            filename="empty.txt",
            file_path=str(empty_file),
            file_size=empty_file.stat().st_size,
            mime_type="text/plain",
            status=FileStatus.PENDING.value,
        )
        session.add(record)
        await session.commit()

    event = {
        "file_id": str(file_id),
        "user_id": str(user_id),
        "file_path": str(empty_file),
        "filename": "empty.txt",
    }

    # Execute worker
    async with session_maker() as session:
        await process_file_event(event=event, db_session=session, redis_client=redis, qdrant_client=qdrant)

    # Verify status transitioned to FAILED with error message
    async with session_maker() as session:
        res = await session.execute(select(DocumentFile).where(DocumentFile.id == file_id))
        doc_record = res.scalar_one()
        assert doc_record.status == FileStatus.FAILED.value
        assert "empty" in doc_record.error_message.lower()


@pytest.mark.asyncio
async def test_file_status_stream_sse_endpoint(test_env: dict):
    """Test GET /api/v1/files/status-stream initiates SSE stream and emits connection event."""
    client: httpx.AsyncClient = test_env["client"]
    token, user_id = await register_and_login(client, "sse_user@example.com")

    # Connect to status stream with query token and max_events=1 for clean exit
    async with client.stream(
        "GET",
        f"/api/v1/files/status-stream?token={token}&max_events=1",
    ) as response:
        assert response.status_code == 200
        assert "text/event-stream" in response.headers.get("content-type", "")

        # Read SSE chunks to completion
        found = False
        async for line in response.aiter_lines():
            if "connected" in line or user_id in line:
                found = True
        assert found
